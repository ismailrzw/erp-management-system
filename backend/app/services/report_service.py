# backend/app/services/report_service.py
"""
Report generation service for Manager exports.
Supports generating group-wise reports as Excel (.xlsx) files via openpyxl.
"""

import io
import re

import openpyxl
from bson import ObjectId
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.extensions import mongo
from app.models.group import COLLECTION as GROUPS_COLLECTION
from app.models.group import Field as GroupField
from app.models.group import Status as GroupStatus
from app.models.user import UserFields


def generate_group_report_excel(
    dept: str | None = None,
    course: str | None = None,
    status: str | None = None,
) -> io.BytesIO:
    """
    Generate a 12-column Excel (.xlsx) workbook report of project groups.

    Columns:
    1. Serial No.
    2. Group Name
    3. Student IDs (Roll Numbers)
    4. Student Names
    5. Project Title
    6. Supervisor
    7. Group Status
    8. Formation Status
    9. Proposal Attached
    10. Core Domain
    11. Department
    12. Course
    """
    query = {GroupField.STATUS: {"$ne": GroupStatus.DELETED}}

    if dept and dept.strip() and dept.lower() != "all":
        query[GroupField.DEPT] = {"$regex": f"^{re.escape(dept.strip())}$", "$options": "i"}

    if course and course.strip() and course.lower() != "all":
        query[GroupField.COURSE] = {"$regex": f"^{re.escape(course.strip())}$", "$options": "i"}

    if status and status.strip() and status.lower() != "all":
        query[GroupField.STATUS] = status.strip().lower()

    groups = list(mongo.db[GROUPS_COLLECTION].find(query).sort(GroupField.NAME, 1))

    # Collect all member OIDs for batch lookup
    all_member_oids = set()
    for g in groups:
        for m in g.get(GroupField.MEMBER_IDS, []):
            if isinstance(m, ObjectId):
                all_member_oids.add(m)
            elif isinstance(m, str) and ObjectId.is_valid(m):
                all_member_oids.add(ObjectId(m))

    user_map = {}
    if all_member_oids:
        users = mongo.db[UserFields.COLLECTION].find(
            {"_id": {"$in": list(all_member_oids)}},
            {UserFields.NAME: 1, UserFields.ROLL: 1},
        )
        for u in users:
            user_map[str(u["_id"])] = u

    # Collect supervisor info if needed
    supervisor_oids = {
        g[GroupField.SUPERVISOR_ID]
        for g in groups
        if g.get(GroupField.SUPERVISOR_ID) and isinstance(g[GroupField.SUPERVISOR_ID], ObjectId)
    }
    supervisor_map = {}
    if supervisor_oids:
        sups = mongo.db[UserFields.COLLECTION].find(
            {"_id": {"$in": list(supervisor_oids)}},
            {UserFields.NAME: 1, UserFields.DOMAINS: 1},
        )
        for s in sups:
            supervisor_map[str(s["_id"])] = s

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Group Report"

    headers = [
        "Serial No.",
        "Group Name",
        "Student IDs",
        "Student Names",
        "Project Title",
        "Supervisor",
        "Group Status",
        "Formation Status",
        "Proposal Attached",
        "Core Domain",
        "Department",
        "Course",
    ]

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, g in enumerate(groups, start=1):
        member_ids = g.get(GroupField.MEMBER_IDS, [])
        rolls = []
        names = []
        for m in member_ids:
            uid = str(m)
            if uid in user_map:
                rolls.append(user_map[uid].get(UserFields.ROLL, ""))
                names.append(user_map[uid].get(UserFields.NAME, ""))
            else:
                rolls.append(uid)
                names.append("")

        student_ids_str = ", ".join(filter(None, rolls))
        student_names_str = ", ".join(filter(None, names))

        sup_id_str = str(g.get(GroupField.SUPERVISOR_ID)) if g.get(GroupField.SUPERVISOR_ID) else None
        sup_name = g.get(GroupField.SUPERVISOR_NAME) or "Not Assigned"

        # Domain determination
        core_domain = "General"
        if sup_id_str and sup_id_str in supervisor_map:
            sup_domains = supervisor_map[sup_id_str].get(UserFields.DOMAINS, [])
            if sup_domains:
                core_domain = ", ".join(sup_domains)

        formation_st = g.get(GroupField.FORMATION_STATUS) or "N/A"
        if formation_st == "on_time":
            formation_label = "On Time"
        elif formation_st == "on_deadline":
            formation_label = "On Deadline"
        elif formation_st == "late":
            formation_label = "Late"
        else:
            formation_label = str(formation_st)

        proposal_attached = "Yes" if g.get(GroupField.PROPOSAL_ATTACHMENT_ID) else "No"
        group_status = (g.get(GroupField.STATUS) or "").capitalize()

        ws.append([
            idx,
            g.get(GroupField.NAME, ""),
            student_ids_str,
            student_names_str,
            g.get(GroupField.PROJECT_TITLE, ""),
            sup_name,
            group_status,
            formation_label,
            proposal_attached,
            core_domain,
            g.get(GroupField.DEPT, ""),
            g.get(GroupField.COURSE, ""),
        ])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
