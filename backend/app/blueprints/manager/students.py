# backend/app/blueprints/manager/students.py
"""
Manager Students API endpoints.

Data lifecycle enforced here:
  1. Raw JSON arrives at the endpoint.
  2. CreateStudentSchema / UpdateStudentSchema validates all fields
     (type, length, no spaces in roll, etc.) and returns a clean dict.
     ValidationError → HTTP 422 with structured ``errors`` payload.
  3. The validated dict is passed to the student service which performs
     business logic (generate email/password, check roll uniqueness) and
     inserts into MongoDB.
  4. Service raises ValueError for constraint violations (409) or returns
     the serialised student dict.
"""

from flask import Blueprint, request
from flask_jwt_extended import get_jwt_identity
from flask_restx import Namespace, Resource, fields, reqparse
from marshmallow import ValidationError
from werkzeug.datastructures import FileStorage

from app.extensions import mongo
from app.models.user import Role
from app.schemas.student_schema import CreateStudentSchema, UpdateStudentSchema
from app.services.bulk_import_service import bulk_create_students, parse_excel
from app.services.student_service import (
    create_student,
    get_student_by_id,
    list_students,
    list_ungrouped_students,
    notify_ungrouped_students,
    permanent_delete_student,
    resend_password_set_email,
    restore_student,
    soft_delete_student,
    update_student,
)

from app.utils.audit import log_audit
from app.utils.decorators import role_required

# ── Blueprint ──────────────────────────────────────────────────
students_bp = Blueprint("manager_students", __name__)

# ── Namespace ──────────────────────────────────────────────────
students_ns = Namespace("manager_students", description="Manager Students operations")

# ── Swagger Models ─────────────────────────────────────────────
student_model = students_ns.model("Student", {
    "name":           fields.String(required=True,  description="Student full name"),
    "roll":           fields.String(required=True,  description="Roll number (no spaces)"),
    "dept":           fields.String(required=True,  description="Department code"),
    "section":        fields.String(required=True,  description="Section"),
    "session":        fields.String(required=True,  description="Academic session"),
    "course":         fields.String(required=True,  description="Course name"),
    "teacher":        fields.String(required=True,  description="Assigned teacher"),
    "recovery_email": fields.String(description="Optional recovery email address"),
})

student_update_model = students_ns.model("StudentUpdate", {
    "name":           fields.String(description="Student full name"),
    "section":        fields.String(description="Section"),
    "course":         fields.String(description="Course name"),
    "teacher":        fields.String(description="Assigned teacher"),
    "recovery_email": fields.String(description="Optional recovery email address"),
})

bulk_import_parser = reqparse.RequestParser()
bulk_import_parser.add_argument(
    "file", location="files", type=FileStorage, required=True,
    help="CSV or XLSX file containing student rows",
)


# ── Routes ──────────────────────────────────────────────────────

@students_ns.route("/")
class StudentList(Resource):
    @students_ns.doc(security="Bearer Auth")
    @role_required(Role.MANAGER)
    def get(self):
        """List all students with pagination and filters."""
        try:
            dept    = request.args.get("dept")
            section = request.args.get("section")
            search  = request.args.get("search", "").strip()
            deleted = request.args.get("deleted", "false").lower() == "true"
            page    = max(1, int(request.args.get("page", 1)))
            limit   = min(100, max(1, int(request.args.get("limit", 20))))

            filters = {
                "dept":    dept,
                "section": section,
                "search":  search,
                "deleted": deleted,
            }

            result = list_students(filters, page, limit)
            return {"success": True, "message": "Students retrieved.", "data": result}, 200

        except Exception as exc:  # noqa: BLE001 - deliberate catch-all, returns error response to client
            return {"success": False, "message": str(exc)}, 500

    @students_ns.doc(security="Bearer Auth")
    @students_ns.expect(student_model)
    @role_required(Role.MANAGER)
    def post(self):
        """Add a new student.  Roll number must be unique and contain no spaces."""
        # Check 1 — schema validation
        try:
            validated = CreateStudentSchema().load(request.get_json() or {})
        except ValidationError as exc:
            return {"success": False, "message": "Validation failed.", "errors": exc.messages}, 422

        # Check 2 — service-layer (DB uniqueness, email generation, bcrypt)
        try:
            result = create_student(validated)
        except ValueError as exc:
            return {"success": False, "message": str(exc)}, 409
        except Exception as exc:  # noqa: BLE001 - deliberate catch-all, returns error response to client
            return {"success": False, "message": str(exc)}, 500

        user_id = get_jwt_identity()
        log_audit(
            mongo.db, user_id, Role.MANAGER, "users", "create",
            target_id=result["student_id"],
            new_value={"roll": validated["roll"]},
        )

        return {"success": True, "message": "Student added successfully.", "data": result}, 201


@students_ns.route("/<student_id>")
class StudentDetail(Resource):
    @students_ns.doc(security="Bearer Auth")
    @role_required(Role.MANAGER)
    def get(self, student_id):
        """Get a single student by ID."""
        try:
            student = get_student_by_id(student_id)
            if not student:
                return {"success": False, "message": "Student not found."}, 404
            return {"success": True, "message": "Student retrieved.", "data": student}, 200
        except Exception as exc:  # noqa: BLE001 - deliberate catch-all, returns error response to client
            return {"success": False, "message": str(exc)}, 500

    @students_ns.doc(security="Bearer Auth")
    @students_ns.expect(student_update_model)
    @role_required(Role.MANAGER)
    def put(self, student_id):
        """Update a student's editable fields (name, section, course, teacher, recovery_email)."""
        # Check 1 — schema validation
        try:
            validated = UpdateStudentSchema().load(request.get_json() or {})
        except ValidationError as exc:
            return {"success": False, "message": "Validation failed.", "errors": exc.messages}, 422

        if not any(v is not None for v in validated.values()):
            return {"success": False, "message": "No fields to update."}, 400

        # Check 2 — service-layer persistence
        try:
            result = update_student(student_id, validated)
        except ValueError as exc:
            return {"success": False, "message": str(exc)}, 404
        except Exception as exc:  # noqa: BLE001 - deliberate catch-all, returns error response to client
            return {"success": False, "message": str(exc)}, 500

        user_id = get_jwt_identity()
        log_audit(mongo.db, user_id, Role.MANAGER, "users", "update", target_id=student_id)

        return {"success": True, "message": "Student updated.", "data": result}, 200

    @students_ns.doc(security="Bearer Auth")
    @role_required(Role.MANAGER)
    def delete(self, student_id):
        """Soft delete a student (moves to recycle bin)."""
        try:
            result = soft_delete_student(student_id)
        except ValueError as exc:
            return {"success": False, "message": str(exc)}, 404
        except Exception as exc:  # noqa: BLE001 - deliberate catch-all, returns error response to client
            return {"success": False, "message": str(exc)}, 500

        user_id = get_jwt_identity()
        log_audit(mongo.db, user_id, Role.MANAGER, "users", "delete", target_id=student_id)

        return {"success": True, "message": "Student deleted.", "data": result}, 200


@students_ns.route("/<student_id>/restore")
class StudentRestore(Resource):
    @students_ns.doc(security="Bearer Auth")
    @role_required(Role.MANAGER)
    def post(self, student_id):
        """Restore a soft-deleted student."""
        try:
            result = restore_student(student_id)
        except ValueError as exc:
            return {"success": False, "message": str(exc)}, 404
        except Exception as exc:  # noqa: BLE001 - deliberate catch-all, returns error response to client
            return {"success": False, "message": str(exc)}, 500

        user_id = get_jwt_identity()
        log_audit(mongo.db, user_id, Role.MANAGER, "users", "restore", target_id=student_id)

        return {"success": True, "message": "Student restored.", "data": result}, 200


@students_ns.route("/<student_id>/permanent")
class StudentPermanentDelete(Resource):
    @students_ns.doc(security="Bearer Auth")
    @role_required(Role.MANAGER)
    def delete(self, student_id):
        """Permanently delete a student (only if already soft-deleted)."""
        try:
            result = permanent_delete_student(student_id)
        except ValueError as exc:
            return {"success": False, "message": str(exc)}, 404
        except Exception as exc:  # noqa: BLE001 - deliberate catch-all, returns error response to client
            return {"success": False, "message": str(exc)}, 500

        user_id = get_jwt_identity()
        log_audit(
            mongo.db, user_id, Role.MANAGER, "users", "permanent_delete",
            target_id=student_id,
        )

        return {"success": True, "message": "Student permanently deleted.", "data": result}, 200


@students_ns.route("/bulk")
@students_ns.route("/bulk-import")
class StudentBulkImport(Resource):
    @students_ns.doc(security="Bearer Auth", consumes=["multipart/form-data"])
    @students_ns.expect(bulk_import_parser)
    @role_required(Role.MANAGER)
    def post(self):
        """Bulk import students from a CSV or XLSX file.  Every row is validated individually."""
        try:
            file = request.files.get("file")
            if not file:
                try:
                    args = bulk_import_parser.parse_args()
                    file = args.get("file")
                except Exception:  # noqa: BLE001
                    file = None
            if not file:
                return {"success": False, "message": "No file uploaded. Please select a CSV or XLSX file."}, 400

            rows, validation_errors = parse_excel(file)
            result = bulk_create_students(rows)
            result["errors"] = validation_errors + result["errors"]
            result["skipped_count"] += len(validation_errors)
            log_audit(
                mongo.db, get_jwt_identity(), Role.MANAGER, "users", "bulk_import",
                new_value={"imported": result["imported_count"], "skipped": result["skipped_count"]},
            )
            return {
                "success": True,
                "message": "Student import completed.",
                "data": {
                    "imported": result["imported_count"],
                    "skipped":  result["skipped_count"],
                    "errors":   result["errors"],
                },
            }, 200
        except ValueError as exc:
            return {"success": False, "message": str(exc)}, 400
        except Exception as exc:
            import logging
            logging.getLogger(__name__).exception("Bulk import exception")
            return {"success": False, "message": f"Unable to import students: {exc!s}"}, 500


@students_ns.route("/<student_id>/resend-password-email")
class StudentResendPasswordEmail(Resource):
    @students_ns.doc(security="Bearer Auth")
    @role_required(Role.MANAGER)
    def post(self, student_id):
        """Regenerate and resend one-time password setup email to student."""
        try:
            result = resend_password_set_email(student_id)
            user_id = get_jwt_identity()
            log_audit(mongo.db, user_id, Role.MANAGER, "users", "resend_password_email", target_id=student_id)
            return {"success": True, "message": "Password setup email has been dispatched.", "data": result}, 200
        except ValueError as exc:
            return {"success": False, "message": str(exc)}, 404
        except Exception as exc:  # noqa: BLE001
            return {"success": False, "message": str(exc)}, 500


@students_ns.route("/ungrouped")
class UngroupedStudentsList(Resource):
    @students_ns.doc(security="Bearer Auth")
    @role_required(Role.MANAGER)
    def get(self):
        """List active students who do not belong to any project group."""
        try:
            dept = request.args.get("dept")
            course = request.args.get("course")
            items = list_ungrouped_students(dept=dept, course=course)
            return {
                "success": True,
                "message": "Ungrouped students retrieved.",
                "data": {
                    "items": items,
                    "total": len(items),
                },
            }, 200
        except Exception as exc:  # noqa: BLE001
            return {"success": False, "message": str(exc)}, 500


@students_ns.route("/ungrouped/export")
class UngroupedStudentsExport(Resource):
    @students_ns.doc(security="Bearer Auth")
    @role_required(Role.MANAGER)
    def get(self):
        """Export ungrouped students as an Excel (.xlsx) spreadsheet."""
        import io
        from datetime import datetime
        from flask import send_file
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        try:
            dept = request.args.get("dept")
            course = request.args.get("course")
            students = list_ungrouped_students(dept=dept, course=course)

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Ungrouped Students"

            headers = ["Serial No.", "Roll Number", "Name", "Department", "Section", "Course", "Session", "Teacher", "Email"]
            ws.append(headers)

            # Styling header
            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style="thin", color="CBD5E1"),
                right=Side(style="thin", color="CBD5E1"),
                top=Side(style="thin", color="CBD5E1"),
                bottom=Side(style="thin", color="CBD5E1"),
            )

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for idx, s in enumerate(students, start=1):
                ws.append([
                    idx,
                    s.get("roll", ""),
                    s.get("name", ""),
                    s.get("dept", ""),
                    s.get("section", ""),
                    s.get("course", ""),
                    s.get("session", ""),
                    s.get("teacher", ""),
                    s.get("email", ""),
                ])

            # Adjust column widths & cell styling
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
                for cell in row:
                    cell.border = thin_border
                    cell.font = Font(name="Calibri", size=10)

            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            filename = f"ungrouped_students_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            return send_file(
                output,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                as_attachment=True,
                download_name=filename,
            )
        except Exception as exc:  # noqa: BLE001
            return {"success": False, "message": str(exc)}, 500


@students_ns.route("/notify-ungrouped")
class NotifyUngroupedStudents(Resource):
    @students_ns.doc(security="Bearer Auth")
    @role_required(Role.MANAGER)
    def post(self):
        """Send reminder email notifications to ungrouped students."""
        try:
            body = request.get_json() or {}
            dept = body.get("dept")
            course = body.get("course")
            custom_message = body.get("message")

            result = notify_ungrouped_students(dept=dept, course=course, custom_message=custom_message)
            user_id = get_jwt_identity()
            log_audit(mongo.db, user_id, Role.MANAGER, "users", "notify_ungrouped", new_value=result)

            return {
                "success": True,
                "message": f"Successfully notified {result['sent_count']} ungrouped student(s).",
                "data": result,
            }, 200
        except Exception as exc:  # noqa: BLE001
            return {"success": False, "message": str(exc)}, 500

